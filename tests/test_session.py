"""M5: multi-file sessions, duplicates across imports, consolidation and exports."""
import csv
import io
from decimal import Decimal

import openpyxl
import pytest

from core import pipeline
from core.dedupe import mark_duplicates
from core.summary import export_csv, export_xlsx
from db import repo
from tests.conftest import CITI_XLSX, citi_profile, citi_sign_rules, dbs_text


@pytest.fixture
def ctx(seeded):
    conn = seeded
    citi_id = repo.create_profile(conn, citi_profile().model_copy(update={"id": None}))
    for rule in citi_sign_rules(citi_id):
        repo.create_rule(conn, rule.model_copy(update={"id": None}))
    return {
        "conn": conn,
        "session_id": repo.create_session(conn, "October 2026"),
        "dbs_id": repo.get_profile_by_name(conn, "DBS Savings CSV").id,
        "citi_id": citi_id,
        "dbs_data": dbs_text().encode("utf-8"),
        "citi_data": CITI_XLSX.read_bytes(),
    }


def import_dbs(ctx, data=None, session_id=None):
    return pipeline.import_file(ctx["conn"], session_id or ctx["session_id"], "dbs_sample.csv",
                                data or ctx["dbs_data"], ctx["dbs_id"])


def import_citi(ctx):
    return pipeline.import_file(ctx["conn"], ctx["session_id"], CITI_XLSX.name, ctx["citi_data"], ctx["citi_id"])


def totals_by_category(summary):
    return {t.name: t.cents for t in summary.totals}


# ---------------------------------------------------------------- duplicates

def test_mark_duplicates():
    class T:
        def __init__(self, h):
            self.row_hash = h
    assert mark_duplicates([T("a"), T("b"), T("c")], {"b"}) == [False, True, False]


def test_reuploading_a_committed_file_adds_zero_rows(ctx):
    conn = ctx["conn"]
    first = import_dbs(ctx)
    pipeline.commit(conn, first)
    second = import_dbs(ctx)

    again = pipeline.file_summary(conn, second)
    assert (again.rows, again.duplicates, again.total_cents) == (8, 8, 0)
    assert all(s.outcome.status == "duplicate" for s in repo.list_transactions(conn, second))
    summary = pipeline.session_summary(conn, ctx["session_id"])
    assert summary.total_cents == 26653  # unchanged
    assert summary.pending == 1  # the refund counted once, not twice


def test_duplicates_only_against_committed_imports(ctx):
    first = import_dbs(ctx)  # still a draft
    second = import_dbs(ctx)
    assert pipeline.file_summary(ctx["conn"], second).duplicates == 0
    pipeline.commit(ctx["conn"], first)  # committing re-checks later imports
    assert pipeline.file_summary(ctx["conn"], second).duplicates == 8


def test_partial_overlap(ctx):
    conn = ctx["conn"]
    pipeline.commit(conn, import_dbs(ctx))
    lines = dbs_text().splitlines()
    overlap = "\n".join(lines[:8] + [  # preamble, header and the first 2 rows of the committed file
        "25 Oct 2026,25 Oct 2026,POS,,,,,,,,,KOUFU HAWKER,6.50,",
        "26 Oct 2026,26 Oct 2026,POS,,,,,,,,,GRAB*FOOD 9911,14.20,",
    ]) + "\n"
    second = import_dbs(ctx, data=overlap.encode())
    summary = pipeline.file_summary(conn, second)
    assert (summary.rows, summary.duplicates) == (4, 2)
    assert summary.total_cents == 650 + 1420


def test_removing_the_original_brings_copies_back(ctx):
    conn = ctx["conn"]
    first = import_dbs(ctx)
    pipeline.commit(conn, first)
    second = import_dbs(ctx)
    pipeline.remove_import(conn, first)
    assert pipeline.file_summary(conn, second).duplicates == 0
    assert pipeline.file_summary(conn, second).total_cents == 26653


def test_earlier_import_keeps_its_rows_when_both_committed(ctx):
    conn = ctx["conn"]
    first = import_dbs(ctx)
    pipeline.commit(conn, first)
    second = import_dbs(ctx)
    pipeline.commit(conn, second)
    pipeline.rerun_session(conn, ctx["session_id"])
    assert pipeline.file_summary(conn, first).duplicates == 0
    assert pipeline.file_summary(conn, second).duplicates == 8


def test_duplicates_across_sessions(ctx):
    conn = ctx["conn"]
    pipeline.commit(conn, import_dbs(ctx))
    november = repo.create_session(conn, "November 2026")
    copy = import_dbs(ctx, session_id=november)
    assert pipeline.file_summary(conn, copy).duplicates == 8


def test_duplicates_survive_rule_changes(ctx):
    conn = ctx["conn"]
    pipeline.commit(conn, import_dbs(ctx))
    second = import_dbs(ctx)
    netflix = next(r for r in repo.list_rules(conn) if r.pattern == "netflix")
    pipeline.save_rule(conn, netflix.model_copy(update={"priority": 1}), ctx["session_id"])
    assert pipeline.file_summary(conn, second).duplicates == 8


# ---------------------------------------------------------------- consolidation

def test_two_files_consolidate_to_sum_of_file_totals(ctx):
    conn = ctx["conn"]
    dbs, citi = import_dbs(ctx), import_citi(ctx)
    file_totals = [pipeline.file_summary(conn, i) for i in (dbs, citi)]
    assert [f.total_cents for f in file_totals] == [26653, 66048]

    summary = pipeline.consolidate(conn, ctx["session_id"])
    assert summary.total_cents == 92701  # 927.01 with the refund pending
    expected: dict[str, int] = {}
    for f in file_totals:
        for t in f.totals:
            expected[t.name] = expected.get(t.name, 0) + t.cents
    assert totals_by_category(summary) == expected
    assert [t.cents for t in summary.totals] == sorted((t.cents for t in summary.totals), reverse=True)
    assert (summary.pending, summary.pending_cents) == (1, 1200)
    assert summary.needs_category == 5  # TOAST BOX + 4 unmatched Citi rows

    assert all(i["status"] == "committed" for i in repo.list_imports(conn, ctx["session_id"]))
    assert repo.get_session(conn, ctx["session_id"])["consolidated_at"] is not None


def test_consolidated_total_follows_decisions(ctx):
    conn = ctx["conn"]
    dbs = import_dbs(ctx)
    import_citi(ctx)
    pipeline.consolidate(conn, ctx["session_id"])
    refund = next(s for s in repo.list_transactions(conn, dbs) if s.txn.description.startswith("REFUND"))
    transport = repo.get_category_by_name(conn, "Transport").id
    pipeline.assign(conn, refund.id, transport)
    summary = pipeline.session_summary(conn, ctx["session_id"])
    assert summary.total_cents == 92701 - 1200
    assert totals_by_category(summary)["Transport"] == 640
    assert summary.pending == 0


def test_summary_query_matches_python_totals(ctx):
    conn = ctx["conn"]
    import_dbs(ctx)
    import_citi(ctx)
    rows = repo.session_category_totals(conn, ctx["session_id"])
    assert {r["category"]: r["cents"] for r in rows} == {
        "Uncategorised": 780 + 64458, "Utilities": 13260, "Groceries": 6425, "Food": 2350 + 1590,
        "Subscriptions": 1998, "Transport": 1840}
    assert rows[0]["category"] == "Uncategorised"


def test_empty_session_summary(ctx):
    summary = pipeline.session_summary(ctx["conn"], ctx["session_id"])
    assert (summary.total_cents, summary.totals, summary.files) == (0, [], [])
    assert summary.share(0) == 0.0


# ---------------------------------------------------------------- exports

@pytest.fixture
def consolidated(ctx):
    import_dbs(ctx)
    import_citi(ctx)
    summary = pipeline.consolidate(ctx["conn"], ctx["session_id"])
    names = {c.id: c.name for c in repo.list_categories(ctx["conn"])}
    return ctx, summary, names


def test_csv_export_opens_as_a_table(consolidated):
    _, summary, _ = consolidated
    data = export_csv(summary)
    assert data.startswith(b"\xef\xbb\xbf")  # BOM: Excel reads it as UTF-8
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0] == ["Category", "Total", "% of total"]
    assert rows[1] == ["Uncategorised", "652.38", "70.4"]
    assert rows[-1] == ["Total", "927.01", "100.0"]
    assert sum(Decimal(r[1]) for r in rows[1:-1]) == Decimal("927.01")


def test_xlsx_export_opens_with_three_sheets(consolidated):
    ctx, summary, names = consolidated
    data = export_xlsx(summary, pipeline.session_transactions(ctx["conn"], ctx["session_id"]), names, "October 2026")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Files", "Transactions"]

    summary_rows = list(wb["Summary"].iter_rows(values_only=True))
    assert summary_rows[0] == ("Category", "Total", "% of total")
    total_row = next(r for r in summary_rows if r[0] == "Total")
    assert total_row[1] == pytest.approx(927.01)
    assert wb["Summary"]["B2"].number_format == "#,##0.00"

    files = list(wb["Files"].iter_rows(values_only=True))
    assert [r[0] for r in files[1:]] == ["dbs_sample.csv", "citi_sample.xlsx"]
    assert [r[-1] for r in files[1:]] == [pytest.approx(266.53), pytest.approx(660.48)]

    txns = list(wb["Transactions"].iter_rows(values_only=True))
    assert len(txns) == 1 + 8 + 6
    assert sum(r[-1] for r in txns[1:]) == pytest.approx(927.01)  # 'Counted' column adds up
    assert txns[1][1].date().isoformat() == "2026-10-01"
