"""Regenerate the fake bank exports in this folder. All data is made up.

    python tests/fixtures/make_fixtures.py

dbs_sample.csv  - DBS savings layout: 5 preamble rows, headers on row 6, description
                  in L, debit in M, credit in N. 8 rows: 7 expenses, 1 refund credit.
                  Expected totals (DESIGN.md section 10): Utilities 132.60, Groceries
                  64.25, Food 23.50, Subscriptions 19.98, Transport 18.40,
                  Uncategorised 7.80 -> file total 266.53 with the refund pending.
citi_sample.xlsx - card layout: spend positive, payments negative, 3 preamble rows,
                  headers on row 4, a 'Total' footer row. Expenses total 660.48.
dbs_with_preamble.csv - real DBS export layout: 5 preamble lines of 1-2 fields,
                  header on row 6 (7 columns), 8 data rows with separate debit and
                  credit columns, 2 footer lines. Same amounts as dbs_sample.csv.
card_negative.xlsx - title row, header on row 2, one amount column where spending is
                  negative and a card payment positive.
"""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent

DBS_HEADERS = [
    "Transaction Date", "Value Date", "Reference", "Transaction Ref1", "Transaction Ref2",
    "Transaction Ref3", "Channel", "Branch", "Teller", "Cheque No", "Status",
    "Description", "Debit Amount", "Credit Amount",
]
# (date, reference, description, debit, credit)
DBS_ROWS = [
    ("01 Oct 2026", "POS", "NTUC FAIRPRICE BEDOK", "64.25", ""),
    ("03 Oct 2026", "POS", "GRAB*FOOD 4821", "23.50", ""),
    ("05 Oct 2026", "POS", "GRAB*TRANSPORT SG", "18.40", ""),
    ("08 Oct 2026", "POS", "NETFLIX.COM", "19.98", ""),
    ("12 Oct 2026", "GIRO", "SP SERVICES LTD", "132.60", ""),
    ("15 Oct 2026", "POS", "TOAST BOX TAMPINES", "7.80", ""),
    ("20 Oct 2026", "ITR", "FUND TRANSFER TO <YOUR NAME>", "1,500.00", ""),
    ("22 Oct 2026", "POS", "REFUND GRAB*TRANSPORT", "", "12.00"),
]

CITI_HEADERS = ["Transaction Date", "Posting Date", "Description", "Amount (SGD)"]
# Mix real Excel dates and text dates to exercise both paths in the normaliser.
CITI_ROWS = [
    (datetime(2026, 10, 2), datetime(2026, 10, 3), "SHOPEE SINGAPORE", 120.00),
    (datetime(2026, 10, 5), datetime(2026, 10, 6), "GRAB*FOOD 8812", 15.90),
    (datetime(2026, 10, 7), datetime(2026, 10, 8), "COLD STORAGE JEWEL", 88.48),
    ("11/10/2026", "12/10/2026", "SINGTEL MOBILE BILL", 45.00),
    ("14/10/2026", "15/10/2026", "SINGAPORE AIRLINES", 391.10),
    ("20/10/2026", "20/10/2026", "PAYMENT - THANK YOU", -500.00),
]


PREAMBLE_HEADERS = [
    "Transaction Date", "Reference", "Debit Amount", "Credit Amount",
    "Transaction Ref1", "Transaction Ref2", "Transaction Ref3",
]
# (date, reference, debit, credit, ref1, ref2)
PREAMBLE_ROWS = [
    ("01 Oct 2026", "POS", "64.25", "", "NTUC FAIRPRICE", "BEDOK"),
    ("03 Oct 2026", "POS", "23.50", "", "GRAB*FOOD", "4821"),
    ("05 Oct 2026", "POS", "18.40", "", "GRAB*TRANSPORT", "SG"),
    ("08 Oct 2026", "UMC", "19.98", "", "NETFLIX.COM", ""),
    ("12 Oct 2026", "GIRO", "132.60", "", "SP SERVICES LTD", ""),
    ("15 Oct 2026", "POS", "7.80", "", "TOAST BOX", "TAMPINES"),
    ("20 Oct 2026", "ITR", "1,500.00", "", "FUND TRANSFER", "TO <YOUR NAME>"),
    ("22 Oct 2026", "ICT", "", "12.00", "REFUND", "GRAB*TRANSPORT"),
]

CARD_ROWS = [
    (datetime(2026, 10, 2), "SHOPEE SINGAPORE", -120.00),
    (datetime(2026, 10, 5), "GRAB*FOOD 8812", -15.90),
    ("07/10/2026", "COLD STORAGE JEWEL", -88.48),
    ("15/10/2026", "PAYMENT RECEIVED - THANK YOU", 300.00),
    ("20/10/2026", "NETFLIX.COM", -19.98),
]


def make_dbs_with_preamble(path: Path) -> None:
    rows = [
        ["Account Details For:", "POSB Savings Account 000-00000-0"],
        ["Statement as at:", "31 Oct 2026"],
        ["Available Balance:", "5,000.00"],
        ["Ledger Balance:", "5,000.00"],
        ["Transaction History"],
        PREAMBLE_HEADERS,
        *[[d, ref, debit, credit, r1, r2, ""] for d, ref, debit, credit, r1, r2 in PREAMBLE_ROWS],
        ["Total Debit Amount:", "1,766.53"],
        ["End of statement"],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def make_card_negative(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append(["Card statement - October 2026"])
    ws.append(["Date", "Description", "Amount"])
    for row in CARD_ROWS:
        ws.append(list(row))
    wb.save(path)


def make_dbs(path: Path) -> None:
    preamble = [
        ["Account Details For:", "POSB Savings Account 000-00000-0"],
        ["Statement as at:", "31 Oct 2026"],
        ["Available Balance:", "5,000.00"],
        ["Ledger Balance:", "5,000.00"],
        [],
    ]
    rows = [DBS_HEADERS]
    for txn_date, ref, desc, debit, credit in DBS_ROWS:
        row = [""] * len(DBS_HEADERS)
        row[0], row[1], row[2], row[11], row[12], row[13] = txn_date, txn_date, ref, desc, debit, credit
        rows.append(row)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(preamble + rows)
        writer.writerow([])  # trailing blank line, as many bank exports have


def make_citi(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Citibank Credit Card Statement"])
    ws.append(["Card ending 0000"])
    ws.append([])
    ws.append(CITI_HEADERS)
    for row in CITI_ROWS:
        ws.append(list(row))
    for cell in ws["A"][4:] + ws["B"][4:]:
        if isinstance(cell.value, datetime):
            cell.number_format = "dd/mm/yyyy"
    ws.append(["Total", "", "", sum(r[3] for r in CITI_ROWS)])
    wb.create_sheet("Notes").append(["Fake data for tests"])
    wb.save(path)


if __name__ == "__main__":
    make_dbs(HERE / "dbs_sample.csv")
    make_citi(HERE / "citi_sample.xlsx")
    make_dbs_with_preamble(HERE / "dbs_with_preamble.csv")
    make_card_negative(HERE / "card_negative.xlsx")
    print(f"Wrote fixtures to {HERE}")
