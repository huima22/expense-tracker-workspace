"""Totals per category for one file and for a whole session, plus the CSV/XLSX exports.

Same semantics as the consolidated query in DESIGN.md section 5: expenses add,
kept income (e.g. a refund kept as Transport) subtracts from its category,
skipped, duplicate and pending-income rows are left out.
"""
from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.models import Category, StoredTxn

UNCATEGORISED = "Uncategorised"


@dataclass
class CategoryTotal:
    name: str
    cents: int
    category_id: int | None = None  # None = Uncategorised
    color: str = "#9AA3AE"


@dataclass
class FileSummary:
    totals: list[CategoryTotal] = field(default_factory=list)  # largest first, Uncategorised included
    total_cents: int = 0
    rows: int = 0
    matched: int = 0  # categorised
    skipped: int = 0
    needs_category: int = 0  # unmatched expenses with no review flag
    needs_category_cents: int = 0
    pending: int = 0  # review-flagged or needs_sign: waiting for a decision
    pending_cents: int = 0
    duplicates: int = 0

    @property
    def ready(self) -> bool:
        """Nothing left to categorise or decide."""
        return self.needs_category == 0 and self.pending == 0


def summarise(stored: Iterable[StoredTxn], categories: Iterable[Category]) -> FileSummary:
    by_id = {c.id: c for c in categories}
    sums: dict[int | None, int] = {}
    s = FileSummary()
    for row in stored:
        o, amount = row.outcome, row.txn.amount_cents
        s.rows += 1
        if o.status == "categorised":
            s.matched += 1
        elif o.status == "skipped":
            s.skipped += 1
        elif o.status == "duplicate":
            s.duplicates += 1
        if o.needs_decision:
            s.pending += 1
            s.pending_cents += amount
        elif o.status == "uncategorised":
            s.needs_category += 1
            s.needs_category_cents += amount
        if o.counts_toward_total:
            key = o.category_id if o.status == "categorised" else None
            sums[key] = sums.get(key, 0) + row.signed_total_cents

    for category_id, cents in sums.items():
        category = by_id.get(category_id)
        if category is None:
            s.totals.append(CategoryTotal(UNCATEGORISED, cents))
        else:
            s.totals.append(CategoryTotal(category.name, cents, category.id, category.color))
    s.totals.sort(key=lambda t: (-t.cents, t.name))
    s.total_cents = sum(sums.values())
    return s


# ---------------------------------------------------------------- session

@dataclass
class FileLine:
    import_id: int
    filename: str
    profile: str
    status: str  # 'draft' or 'committed'
    summary: FileSummary


@dataclass
class SessionSummary:
    totals: list[CategoryTotal] = field(default_factory=list)  # from the consolidated query, largest first
    total_cents: int = 0
    files: list[FileLine] = field(default_factory=list)
    needs_category: int = 0
    needs_category_cents: int = 0
    pending: int = 0
    pending_cents: int = 0

    def share(self, cents: int) -> float:
        """A category's fraction of the grand total (0 when the total isn't positive)."""
        return cents / self.total_cents if self.total_cents > 0 else 0.0


def build_session_summary(category_rows: Iterable[dict], files: list[FileLine]) -> SessionSummary:
    """Combine the summary query's rows (session_category_totals) with per-file summaries.

    The grand total is the sum of the query's rows, as DESIGN.md section 5 says.
    """
    totals = [CategoryTotal(r["category"], r["cents"], r["category_id"], r["color"]) for r in category_rows]
    return SessionSummary(
        totals=totals,
        total_cents=sum(t.cents for t in totals),
        files=files,
        needs_category=sum(f.summary.needs_category for f in files),
        needs_category_cents=sum(f.summary.needs_category_cents for f in files),
        pending=sum(f.summary.pending for f in files),
        pending_cents=sum(f.summary.pending_cents for f in files),
    )


# ---------------------------------------------------------------- exports

def _amount(cents: int) -> Decimal:
    return Decimal(cents) / 100


def _summary_rows(s: SessionSummary) -> list[list]:
    rows = [[t.name, _amount(t.cents), round(s.share(t.cents) * 100, 1)] for t in s.totals]
    rows.append(["Total", _amount(s.total_cents), 100.0 if s.total_cents > 0 else 0.0])
    return rows


def export_csv(s: SessionSummary) -> bytes:
    """Category, total and share, with a Total row. UTF-8 with BOM so Excel opens it cleanly."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["Category", "Total", "% of total"])
    for name, amount, pct in _summary_rows(s):
        writer.writerow([name, f"{amount:.2f}", f"{pct:.1f}"])
    return out.getvalue().encode("utf-8-sig")


_TYPE = {True: "Expense", False: "Income", None: ""}


def export_xlsx(s: SessionSummary, transactions: Iterable[tuple[str, StoredTxn]],
                category_names: dict[int, str], session_name: str = "") -> bytes:
    """Workbook with Summary, Files and Transactions sheets; amounts as numbers."""
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(ws, headers: list[str], rows: Iterable[list], money_cols: tuple[int, ...] = ()) -> None:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(row)
        for col in money_cols:
            for cell in ws[get_column_letter(col)][1:]:
                cell.number_format = "#,##0.00"
        for i, header in enumerate(headers, start=1):
            width = max([len(str(header))] + [len(str(c.value or "")) for c in ws[get_column_letter(i)]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"

    summary_ws = wb.active
    summary_ws.title = "Summary"
    sheet(summary_ws, ["Category", "Total", "% of total"], _summary_rows(s), money_cols=(2,))
    for cell in summary_ws[summary_ws.max_row]:
        cell.font = bold
    if session_name:
        summary_ws.append([])
        summary_ws.append([f"Session: {session_name}"])

    sheet(wb.create_sheet("Files"),
          ["File", "Profile", "Status", "Rows", "Skipped", "Duplicates", "Needs category", "Needs decision", "Total"],
          ([f.filename, f.profile, f.status, f.summary.rows, f.summary.skipped, f.summary.duplicates,
            f.summary.needs_category, f.summary.pending, _amount(f.summary.total_cents)] for f in s.files),
          money_cols=(9,))

    def txn_row(filename: str, row: StoredTxn) -> list:
        o, t = row.outcome, row.txn
        category = category_names.get(o.category_id, "") if o.status == "categorised" else ""
        return [filename, t.txn_date, t.description, _amount(t.signed_amount_cents), _TYPE[o.is_debit],
                o.status, category, o.review_reason or "", _amount(row.signed_total_cents)]

    tx_ws = wb.create_sheet("Transactions")
    sheet(tx_ws, ["File", "Date", "Description", "Amount (as in file)", "Type", "Status", "Category",
                  "Review reason", "Counted"],
          (txn_row(f, r) for f, r in transactions), money_cols=(4, 9))
    for cell in tx_ws["B"][1:]:
        cell.number_format = "yyyy-mm-dd"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
