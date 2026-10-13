"""XLSX parser (openpyxl, values only). Converts typed cells to text so the frame is all strings.

Date cells become ISO text ('2026-10-03'); the normaliser accepts ISO dates in
addition to the profile's date_format. Numbers become their shortest exact text.
"""
from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, time

import openpyxl
import pandas as pd
from openpyxl.utils.exceptions import InvalidFileException

from core.parsers.base import ParseError, rows_to_frame


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


def _open(data: bytes):
    try:
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise ParseError(f"not a readable .xlsx file: {exc}") from exc


class XlsxParser:
    extensions = (".xlsx",)

    def read(self, data: bytes, *, sheet_name: str | None = None, encoding: str | None = None) -> pd.DataFrame:
        wb = _open(data)
        try:
            if sheet_name is None:
                ws = wb.worksheets[0]
            elif sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                raise ParseError(f"sheet {sheet_name!r} not found; sheets: {', '.join(wb.sheetnames)}")
            rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        return rows_to_frame(rows)

    def sheet_names(self, data: bytes) -> list[str]:
        wb = _open(data)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
